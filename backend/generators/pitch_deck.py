"""
Pitch deck generator using python-pptx.
Builds a 7-slide deck with a Midnight Executive palette.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# Midnight Executive palette
NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
GOLD = RGBColor(0xC5, 0xA5, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x33)
GRAY = RGBColor(0x77, 0x77, 0x88)
GREEN = RGBColor(0x0F, 0x8B, 0x5C)
RED = RGBColor(0xB3, 0x3A, 0x3A)


def _set_text(tf, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, font="Calibri"):
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_text_box(slide, x, y, w, h, text, **kwargs):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _set_text(box.text_frame, text, **kwargs)
    box.text_frame.word_wrap = True
    return box


def _add_filled_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _fmt_money(v) -> str:
    if v is None:
        return "N/A"
    try:
        v = float(v)
    except (TypeError, ValueError):
        return "N/A"
    abs_v = abs(v)
    if abs_v >= 1e12:
        return f"${v / 1e12:.2f}T"
    if abs_v >= 1e9:
        return f"${v / 1e9:.2f}B"
    if abs_v >= 1e6:
        return f"${v / 1e6:.2f}M"
    return f"${v:,.0f}"


def _bullets(text: str, limit: int = 4) -> list[str]:
    if not text:
        return []
    out = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        for prefix in ("- ", "* ", "• "):
            if line.startswith(prefix):
                line = line[len(prefix):]
                break
        if line and line[0].isdigit() and len(line) > 2 and line[1] in ".)":
            line = line[2:].strip()
        if line:
            out.append(line)
        if len(out) >= limit:
            break
    return out


def generate_pitch_deck(state: dict) -> bytes:
    """Build pitch deck and return bytes."""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    profile = state["company_data"]["profile"]
    metrics = state["company_data"]["metrics"]
    val = state.get("dcf_valuation", {})
    rec = state.get("recommendation", "HOLD")
    target = state.get("target_price", 0)
    current = metrics.get("current_price") or 0
    upside = val.get("upside_pct", 0)
    rec_color = {"BUY": GREEN, "HOLD": GOLD, "SELL": RED}.get(rec, GRAY)

    # =========================================================
    # SLIDE 1: Title (dark)
    # =========================================================
    slide = prs.slides.add_slide(blank)
    _add_filled_rect(slide, 0, 0, 13.33, 7.5, NAVY)
    # Gold accent block
    _add_filled_rect(slide, 0.7, 3.0, 0.15, 1.5, GOLD)
    _add_text_box(slide, 1.0, 2.85, 11, 1.0,
                  f"{profile['name']}", size=44, bold=True, color=WHITE, font="Georgia")
    _add_text_box(slide, 1.0, 3.85, 11, 0.5,
                  f"({profile['ticker']}) — {profile.get('sector', 'N/A')}",
                  size=20, color=ICE, font="Calibri")
    _add_text_box(slide, 1.0, 4.55, 11, 0.4,
                  "Equity Research & Investment Thesis",
                  size=14, color=GOLD, font="Calibri")
    _add_text_box(slide, 1.0, 6.7, 11, 0.4,
                  f"Generated {datetime.now().strftime('%B %d, %Y')}",
                  size=11, color=ICE, font="Calibri")

    # =========================================================
    # SLIDE 2: Recommendation (light)
    # =========================================================
    slide = prs.slides.add_slide(blank)
    _add_text_box(slide, 0.6, 0.4, 12, 0.6,
                  "Recommendation", size=32, bold=True, color=NAVY, font="Georgia")
    _add_filled_rect(slide, 0.6, 1.1, 12, 0.03, GOLD)

    # Big recommendation badge
    _add_filled_rect(slide, 0.6, 1.6, 4.0, 2.0, rec_color)
    _add_text_box(slide, 0.6, 1.95, 4.0, 0.5,
                  "RATING", size=14, color=WHITE, align=PP_ALIGN.CENTER, font="Calibri")
    _add_text_box(slide, 0.6, 2.4, 4.0, 1.0,
                  rec, size=72, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, font="Georgia")

    # Right-side stats grid
    stats = [
        ("Target Price", f"${target:,.2f}"),
        ("Current Price", f"${current:,.2f}"),
        ("Implied Upside", f"{upside * 100:+.1f}%"),
        ("Market Cap", _fmt_money(profile.get("market_cap"))),
    ]
    sx = 5.0
    for i, (label, value) in enumerate(stats):
        row = i // 2
        col = i % 2
        x = sx + col * 3.9
        y = 1.6 + row * 1.05
        _add_filled_rect(slide, x, y, 3.7, 0.9, ICE)
        _add_text_box(slide, x + 0.15, y + 0.08, 3.4, 0.3,
                      label, size=11, color=GRAY, font="Calibri")
        _add_text_box(slide, x + 0.15, y + 0.38, 3.4, 0.5,
                      value, size=22, bold=True, color=NAVY, font="Georgia")

    # Exec summary preview
    summary = (state.get("executive_summary") or "").split("\n\n")[0][:500]
    _add_text_box(slide, 0.6, 4.0, 12, 0.4,
                  "Summary", size=16, bold=True, color=NAVY, font="Georgia")
    _add_text_box(slide, 0.6, 4.45, 12, 2.5,
                  summary, size=12, color=DARK, font="Calibri")

    # =========================================================
    # SLIDE 3: Business overview
    # =========================================================
    slide = prs.slides.add_slide(blank)
    _add_text_box(slide, 0.6, 0.4, 12, 0.6,
                  "Business Overview", size=32, bold=True, color=NAVY, font="Georgia")
    _add_filled_rect(slide, 0.6, 1.1, 12, 0.03, GOLD)

    # Left col: text
    overview = state.get("business_overview", "") or ""
    _add_text_box(slide, 0.6, 1.5, 7.5, 5.5,
                  overview, size=13, color=DARK, font="Calibri")

    # Right col: facts box
    _add_filled_rect(slide, 8.5, 1.5, 4.3, 5.5, ICE)
    _add_text_box(slide, 8.7, 1.7, 4.0, 0.4,
                  "At a Glance", size=14, bold=True, color=NAVY, font="Georgia")
    facts = [
        ("Sector", profile.get("sector", "N/A")),
        ("Industry", profile.get("industry", "N/A")),
        ("Country", profile.get("country", "N/A")),
        ("Employees", f"{profile.get('employees'):,}" if profile.get("employees") else "N/A"),
        ("Exchange", profile.get("exchange", "N/A")),
    ]
    fy = 2.2
    for label, value in facts:
        _add_text_box(slide, 8.7, fy, 4.0, 0.3,
                      label.upper(), size=9, color=GRAY, font="Calibri")
        _add_text_box(slide, 8.7, fy + 0.27, 4.0, 0.4,
                      str(value)[:40], size=13, bold=True, color=DARK, font="Calibri")
        fy += 0.85

    # =========================================================
    # SLIDE 4: Investment thesis (Bull case)
    # =========================================================
    slide = prs.slides.add_slide(blank)
    _add_text_box(slide, 0.6, 0.4, 12, 0.6,
                  "Investment Thesis", size=32, bold=True, color=NAVY, font="Georgia")
    _add_filled_rect(slide, 0.6, 1.1, 12, 0.03, GOLD)
    _add_text_box(slide, 0.6, 1.2, 12, 0.4,
                  "Why we believe in this company", size=14, color=GRAY, font="Calibri")

    thesis_bullets = _bullets(state.get("investment_thesis", ""), limit=4)
    by = 1.9
    for i, b in enumerate(thesis_bullets, 1):
        # Number circle
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.6), Inches(by), Inches(0.6), Inches(0.6)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = GREEN
        circle.line.fill.background()
        _set_text(circle.text_frame, str(i),
                  size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        circle.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        # Bullet text
        _add_text_box(slide, 1.4, by + 0.05, 11.4, 1.15,
                      b, size=14, color=DARK, font="Calibri")
        by += 1.25

    # =========================================================
    # SLIDE 5: Risks (Bear case)
    # =========================================================
    slide = prs.slides.add_slide(blank)
    _add_text_box(slide, 0.6, 0.4, 12, 0.6,
                  "Key Risks", size=32, bold=True, color=NAVY, font="Georgia")
    _add_filled_rect(slide, 0.6, 1.1, 12, 0.03, GOLD)
    _add_text_box(slide, 0.6, 1.2, 12, 0.4,
                  "What could derail the thesis", size=14, color=GRAY, font="Calibri")

    risk_bullets = _bullets(state.get("risks", ""), limit=4)
    ry = 1.9
    for i, b in enumerate(risk_bullets, 1):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.6), Inches(ry), Inches(0.6), Inches(0.6)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = RED
        circle.line.fill.background()
        _set_text(circle.text_frame, "!",
                  size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        circle.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _add_text_box(slide, 1.4, ry + 0.05, 11.4, 1.15,
                      b, size=14, color=DARK, font="Calibri")
        ry += 1.25

    # =========================================================
    # SLIDE 6: DCF valuation
    # =========================================================
    slide = prs.slides.add_slide(blank)
    _add_text_box(slide, 0.6, 0.4, 12, 0.6,
                  "DCF Valuation", size=32, bold=True, color=NAVY, font="Georgia")
    _add_filled_rect(slide, 0.6, 1.1, 12, 0.03, GOLD)

    assumptions = state.get("dcf_assumptions", {})
    # Assumptions strip
    assump_text = (
        f"WACC: {assumptions.get('wacc', 0) * 100:.1f}%   |   "
        f"Terminal Growth: {assumptions.get('terminal_growth', 0) * 100:.1f}%   |   "
        f"FCF Margin: {assumptions.get('fcf_margin', 0) * 100:.1f}%"
    )
    _add_text_box(slide, 0.6, 1.3, 12, 0.4,
                  assump_text, size=13, color=GRAY, font="Calibri")

    # Projected FCF table
    rev = val.get("projected_revenues", [])
    fcf = val.get("projected_fcfs", [])
    pv = val.get("pv_fcfs", [])

    if rev:
        n = len(rev)
        col_w = 9.5 / (n + 1)
        # Header row
        _add_filled_rect(slide, 0.6, 2.0, col_w, 0.5, NAVY)
        _add_text_box(slide, 0.6, 2.05, col_w, 0.4,
                      "($M)", size=12, bold=True, color=WHITE,
                      align=PP_ALIGN.CENTER, font="Calibri")
        for i in range(n):
            _add_filled_rect(slide, 0.6 + col_w * (i + 1), 2.0, col_w, 0.5, NAVY)
            _add_text_box(slide, 0.6 + col_w * (i + 1), 2.05, col_w, 0.4,
                          f"Y{i + 1}", size=12, bold=True, color=WHITE,
                          align=PP_ALIGN.CENTER, font="Calibri")
        # Data rows
        rows = [("Revenue", rev), ("FCF", fcf), ("PV(FCF)", pv)]
        for row_idx, (label, values) in enumerate(rows):
            y = 2.5 + row_idx * 0.5
            bg = ICE if row_idx % 2 == 0 else WHITE
            for i in range(n + 1):
                _add_filled_rect(slide, 0.6 + col_w * i, y, col_w, 0.5, bg)
            _add_text_box(slide, 0.7, y + 0.08, col_w - 0.1, 0.4,
                          label, size=11, bold=True, color=NAVY, font="Calibri")
            for i, v in enumerate(values):
                _add_text_box(slide, 0.6 + col_w * (i + 1), y + 0.08, col_w, 0.4,
                              _fmt_money(v / 1e6), size=11, color=DARK,
                              align=PP_ALIGN.CENTER, font="Calibri")

    # Big implied price callout (right side)
    _add_filled_rect(slide, 10.3, 2.0, 2.4, 2.5, rec_color)
    _add_text_box(slide, 10.3, 2.15, 2.4, 0.4,
                  "IMPLIED", size=11, color=WHITE, align=PP_ALIGN.CENTER, font="Calibri")
    _add_text_box(slide, 10.3, 2.5, 2.4, 0.4,
                  "PRICE", size=11, color=WHITE, align=PP_ALIGN.CENTER, font="Calibri")
    _add_text_box(slide, 10.3, 3.0, 2.4, 0.9,
                  f"${target:,.0f}", size=36, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, font="Georgia")
    _add_text_box(slide, 10.3, 4.0, 2.4, 0.4,
                  f"{upside * 100:+.1f}% upside", size=12, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, font="Calibri")

    # Rationale
    _add_text_box(slide, 0.6, 5.0, 12, 0.4,
                  "Rationale", size=14, bold=True, color=NAVY, font="Georgia")
    _add_text_box(slide, 0.6, 5.4, 12, 1.8,
                  assumptions.get("rationale", "N/A"),
                  size=12, color=DARK, font="Calibri")

    # =========================================================
    # SLIDE 7: Closing (dark)
    # =========================================================
    slide = prs.slides.add_slide(blank)
    _add_filled_rect(slide, 0, 0, 13.33, 7.5, NAVY)
    _add_filled_rect(slide, 0.7, 3.1, 0.15, 1.3, GOLD)
    _add_text_box(slide, 1.0, 3.0, 11, 0.8,
                  rec, size=64, bold=True, color=rec_color, font="Georgia")
    _add_text_box(slide, 1.0, 3.85, 11, 0.5,
                  f"Target Price: ${target:,.2f}  |  {upside * 100:+.1f}% upside",
                  size=20, color=WHITE, font="Calibri")
    _add_text_box(slide, 1.0, 4.5, 11, 0.4,
                  f"{profile['name']} ({profile['ticker']})",
                  size=14, color=ICE, font="Calibri")
    _add_text_box(slide, 1.0, 6.7, 11, 0.4,
                  "Generated by AI Equity Research Agent  ·  Not investment advice",
                  size=10, color=GRAY, font="Calibri")

    # Save to bytes
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
